#!/usr/bin/env python3
"""ADAM-3600 tag import builder (reverse-engineered from adam-3600_tags_template.xls).

Subcommands
  template OUT.xlsx                         blank canonical input template (Tags + Reference sheets)
  build IN --base {0,1} -o OUT.xls [...]    canonical input (.xlsx/.xls/.csv) -> ADAM-3600 import .xls
  decode EXPORT.xls -o OUT.xlsx             ADAM-3600 export -> canonical input (addresses written base-1)
  selftest                                  rebuild the template's VA_Temp rows and compare cell by cell

build exit codes
  0  output written and re-verified
  1  errors - nothing written; the user must fix the listed rows
  2  confirmations needed - nothing written; rerun with --confirm only after the user approves them
"""
import argparse, csv, math, os, re, sys

# ---------------------------------------------------------------- template facts
HEADERS = ['name', 'type', 'initialValue', 'dataType', 'address', 'Scale.spanLow', 'Scale.spanHigh',
           'signalReverse', 'description', 'scalingType', 'Scale.scale', 'Scale.bias', 'Scale.outputLow',
           'Scale.outputHigh', 'Scale.clampToLow', 'Scale.clampToHigh', 'Scale.clampToZero', 'readWrite',
           'template', 'length', 'conversionType', 'startBit', 'scanRate']
USER_TAG_HEADERS = ['name', 'type', 'initialValue', 'dataType', 'description', 'readWrite', 'retain']
CALC_TAG_HEADERS = ['name', 'type', 'initialValue', 'period', 'description', 'expression', 'dataType',
                    'readWrite'] + ['parameters' + c for c in 'ABCDEFGH']
WIDTHS = {
    'device':   [7168, 1216, 2528, 2208, 1952, 3360, 3456, 3168, 22304, 2720, 2624, 2400, 3648, 3744, 4096,
                 4192, 4192, 2304, 2080, 1568, 3488, 1760, 2240],
    'BoardIO':  [3168, 1216, 2528, 2208, 1952, 3360, 9920, 3168, 2496, 2720, 2624, 2400, 3648, 3744, 4096,
                 4192, 4192, 2304, 2080, 1568, 3488, 1760, 2240],
    'User Tag': [1472, 1216, 2528, 2208, 2496, 2304, 1472],
    'Calc Tag': [1472, 1216, 2528, 1600, 2496, 2496, 2208, 2304, 2912, 2912, 2912, 2912, 2912],
}
RESERVED_SHEETS = {'boardio', 'user tag', 'calc tag'}
MAX_NAME = 26          # full "Device:Tag" length limit
FLT_MAX = 3.4028234663852886e+38

# register area: address prefix digit, 'template' column value
AREAS = {'CS': ('0', 0.0), 'IS': ('1', 1.0), 'IR': ('3', 2.0), 'HR': ('4', 3.0)}
PREFIX_TO_AREA = {v[0]: k for k, v in AREAS.items()}
AREA_ALIASES = {}
for _k, _names in {
    'CS': ['CS', 'COIL', 'COILS', 'COIL STATUS', 'FC01', 'FC1', '0', '0X'],
    'IS': ['IS', 'INPUT STATUS', 'DISCRETE INPUT', 'DISCRETE INPUTS', 'FC02', 'FC2', '1', '1X'],
    'IR': ['IR', 'INPUT REGISTER', 'INPUT REGISTERS', 'FC04', 'FC4', '3', '3X'],
    'HR': ['HR', 'HOLDING REGISTER', 'HOLDING REGISTERS', 'FC03', 'FC3', '4', '4X'],
}.items():
    for _n in _names:
        AREA_ALIASES[_n] = _k

# format name -> (conversionType, meaning). Length is independent: it sets how many bits are read.
FORMATS = {
    'DISCRETE':    (21,  'Discrete bit (coil, input status, or one bit of a register)'),
    'UINT':        (21,  'Unsigned integer'),
    'INT':         (22,  'Signed integer'),
    'UINT_SWAP':   (1,   'Unsigned integer, bytes swapped'),
    'INT_SWAP':    (2,   'Signed integer, bytes swapped'),
    'HIWORD_INT':  (16,  'High-word signed integer'),
    'HIWORD_UINT': (17,  'High-word unsigned integer'),
    'FLOAT':       (25,  'IEEE 754 floating point'),
    'RFLOAT':      (26,  'IEEE 754 floating point, reversed word order'),
    'BCD':         (35,  'BCD'),
    'DOUBLE':      (61,  'Double precision (plain double)'),
    'DOUBLE_BE':   (150, 'Double precision, big-endian (IEEE)'),
    'DOUBLE_LE':   (151, 'Double precision, little-endian'),
}
CONV_TO_ANALOG = {v[0]: k for k, v in FORMATS.items() if k != 'DISCRETE'}
EXPECTED_LEN = {'FLOAT': 32, 'RFLOAT': 32, 'DOUBLE': 64, 'DOUBLE_BE': 64, 'DOUBLE_LE': 64}
ANALOG_LENGTHS = (8, 16, 32, 64)

IN_COLS = ['Device', 'Tag', 'RegisterType', 'Address', 'Format', 'Length', 'Access', 'StartBit',
           'ActiveLow', 'ScanRate', 'Description']
REQUIRED_COLS = ['Device', 'Tag', 'RegisterType', 'Address', 'Format', 'Access', 'ScanRate']


# ---------------------------------------------------------------- helpers
def key(h):
    return re.sub(r'[\s_\-./]', '', str(h)).lower()


def cell_str(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_int(s):
    s = s.strip()
    if re.fullmatch(r'[+-]?\d+', s):
        return int(s)
    if re.fullmatch(r'0[xX][0-9a-fA-F]+', s):
        return int(s, 16)
    f = float(s)                      # raises ValueError if not numeric
    if not f.is_integer():
        raise ValueError('not a whole number')
    return int(f)


def parse_format(raw):
    """'INT32' -> ('INT', 32); 'INT16_SWAP' -> ('INT_SWAP', 16); 'FLOAT' -> ('FLOAT', None)."""
    s = re.sub(r'[\s\-]+', '_', raw.strip().upper())
    if s in FORMATS:
        return s, None
    m = re.fullmatch(r'([A-Z_]*?)(8|16|32|64)([A-Z_]*)', s)
    if m:
        base = re.sub(r'_+', '_', (m.group(1) + m.group(3))).strip('_')
        if base in FORMATS:
            return base, int(m.group(2))
    return None, None


def read_table(path, sheet=None):
    ext = os.path.splitext(path)[1].lower()
    rows = []
    if ext == '.csv':
        with open(path, newline='', encoding='utf-8-sig') as f:
            data = list(csv.reader(f))
    elif ext in ('.xlsx', '.xlsm'):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else (wb['Tags'] if 'Tags' in wb.sheetnames else wb.worksheets[0])
        data = [list(r) for r in ws.iter_rows(values_only=True)]
    elif ext == '.xls':
        import xlrd
        wb = xlrd.open_workbook(path)
        ws = wb.sheet_by_name(sheet) if sheet else (
            wb.sheet_by_name('Tags') if 'Tags' in wb.sheet_names() else wb.sheet_by_index(0))
        data = [ws.row_values(r) for r in range(ws.nrows)]
    else:
        sys.exit('Unsupported input type %s (use .xlsx, .xls or .csv)' % ext)
    if not data:
        sys.exit('Input sheet is empty')
    hdr = [key(h) for h in data[0]]
    colmap = {}
    for c in IN_COLS:
        if key(c) in hdr:
            colmap[c] = hdr.index(key(c))
    missing = [c for c in REQUIRED_COLS if c not in colmap]
    if missing:
        sys.exit('Input is not in canonical form - missing columns: %s. Map it to the canonical '
                 'columns first (%s).' % (', '.join(missing), ', '.join(IN_COLS)))
    for i, r in enumerate(data[1:], start=2):
        rec = {}
        for c in IN_COLS:
            v = r[colmap[c]] if c in colmap and colmap[c] < len(r) else None
            if c == 'Description':
                rec[c] = '' if v is None else (str(v) if not isinstance(v, float) else cell_str(v))
            else:
                rec[c] = cell_str(v)
        if any(rec[c] for c in IN_COLS):
            rows.append((i, rec))
    return rows


# ---------------------------------------------------------------- validation + row building
def build_tags(rows, base):
    errors, confirms, notes = [], [], []
    tags = []
    names_seen, names_ci = {}, {}
    dev_ci = {}

    for rn, r in rows:
        E = lambda m: errors.append((rn, m))
        C = lambda m: confirms.append((rn, m))
        n_err = len(errors)

        dev, tag = r['Device'], r['Tag']
        for label, val in (('Device', dev), ('Tag', tag)):
            if not val:
                E('%s is empty' % label)
            elif ':' in val:
                E('%s "%s" contains ":" (the Device:Tag separator)' % (label, val))
        if dev:
            if dev.lower() in RESERVED_SHEETS:
                E('Device name "%s" is reserved (BoardIO / User Tag / Calc Tag)' % dev)
            if re.search(r'[\[\]*?/\\]', dev):
                E('Device name "%s" contains a character Excel forbids in sheet names ([ ] * ? / \\)' % dev)
            prev = dev_ci.setdefault(dev.lower(), dev)
            if prev != dev:
                E('Device "%s" differs from device "%s" only by letter case (sheet names would clash)'
                  % (dev, prev))
        full = '%s:%s' % (dev, tag)
        if dev and tag:
            if len(full) > MAX_NAME:
                E('Name "%s" is %d characters; the limit is %d' % (full, len(full), MAX_NAME))
            odd = sorted(set(re.sub(r'[A-Za-z0-9_.:]', '', dev + tag)))
            if odd:
                C('Name "%s" contains unusual characters: %s' % (full, ' '.join(repr(c) for c in odd)))
            if full in names_seen:
                E('Duplicate name "%s" (also row %d)' % (full, names_seen[full]))
            else:
                names_seen[full] = rn
                prev = names_ci.setdefault(full.lower(), (full, rn))
                if prev[0] != full:
                    C('Name "%s" differs from "%s" (row %d) only by letter case' % (full, prev[0], prev[1]))

        # register type
        area = AREA_ALIASES.get(re.sub(r'\s+', ' ', r['RegisterType'].upper()).strip())
        if not r['RegisterType']:
            E('RegisterType is empty')
        elif not area:
            E('RegisterType "%s" not recognised (use CS, IS, IR or HR)' % r['RegisterType'])

        # address
        addr_out = None
        if not r['Address']:
            E('Address is empty')
        else:
            try:
                a = to_int(r['Address'])
                addr_out = a + 1 if base == 0 else a
                if not 1 <= addr_out <= 9999:
                    hint = ''
                    if a > 9999:
                        hint = ' - it may already include the register-type prefix (e.g. 40001); give the number inside the area'
                    E('Address %s (base-%d) gives %d, outside 1..9999%s' % (r['Address'], base, addr_out, hint))
                    addr_out = None
            except ValueError:
                E('Address "%s" is not a whole number' % r['Address'])

        # format + length
        fmt, sfx_len = parse_format(r['Format']) if r['Format'] else (None, None)
        if not r['Format']:
            E('Format is empty')
        elif not fmt:
            E('Format "%s" not recognised (use one of: %s, optionally with a width suffix like INT32)'
              % (r['Format'], ', '.join(FORMATS)))
        length = None
        col_len = None
        if r['Length']:
            try:
                col_len = to_int(r['Length'])
            except ValueError:
                E('Length "%s" is not a whole number' % r['Length'])
        if fmt:
            if sfx_len and col_len and sfx_len != col_len:
                E('Format "%s" implies %d bits but Length is %d' % (r['Format'], sfx_len, col_len))
            length = col_len or sfx_len
            if fmt == 'DISCRETE':
                if length is None:
                    length = 1
                elif length != 1:
                    C('DISCRETE tag with Length %d (discrete tags are normally 1)' % length)
            else:
                if length is None:
                    E('Length is missing for %s (give 8, 16, 32 or 64%s)' % (
                        fmt, ', normally %d' % EXPECTED_LEN[fmt] if fmt in EXPECTED_LEN else ''))
                elif length not in ANALOG_LENGTHS:
                    E('Length %d not allowed (8, 16, 32 or 64)' % length)
                    length = None
                else:
                    exp = EXPECTED_LEN.get(fmt)
                    if exp and length != exp:
                        C('%s with Length %d (normally %d)' % (fmt, length, exp))
                    if fmt == 'BCD' and length == 64:
                        C('BCD with Length 64')
            if area in ('CS', 'IS') and fmt != 'DISCRETE':
                E('%s registers hold single bits; Format must be DISCRETE (got %s)' % (area, fmt))

        # access
        acc = re.sub(r'[\s/_\-]', '', r['Access'].upper())
        rw = {'R': '1', 'RO': '1', 'READ': '1', 'READONLY': '1', '1': '1',
              'RW': '3', 'READWRITE': '3', '3': '3'}.get(acc)
        if not r['Access']:
            E('Access is empty (R or RW)')
        elif not rw:
            E('Access "%s" not recognised (R or RW)' % r['Access'])
        elif rw == '3' and area in ('IS', 'IR'):
            E('%s is read-only by nature but Access is RW' % area)

        # start bit
        start = 0
        if r['StartBit']:
            try:
                start = to_int(r['StartBit'])
            except ValueError:
                E('StartBit "%s" is not a whole number' % r['StartBit'])
        if fmt == 'DISCRETE' and area in ('HR', 'IR'):
            if not r['StartBit']:
                E('StartBit is required for a discrete bit on %s (0..15)' % area)
            elif not 0 <= start <= 15:
                E('StartBit %d outside 0..15' % start)
        elif area in ('CS', 'IS'):
            if start != 0:
                E('StartBit must be blank or 0 on %s' % area)
        elif fmt and fmt != 'DISCRETE':
            if not 0 <= start <= 15:
                E('StartBit %d outside 0..15' % start)
            elif start != 0:
                C('%s tag starts at bit %d (every analog tag in the template starts at bit 0)' % (fmt, start))

        # active low
        al = r['ActiveLow'].upper()
        rev = {'': 'false', 'FALSE': 'false', 'NO': 'false', 'N': 'false', '0': 'false',
               'TRUE': 'true', 'YES': 'true', 'Y': 'true', '1': 'true'}.get(al)
        if rev is None:
            E('ActiveLow "%s" not recognised (TRUE/FALSE)' % r['ActiveLow'])
        elif rev == 'true' and fmt and fmt != 'DISCRETE':
            C('ActiveLow set on a %s tag (only discrete tags use it in the template)' % fmt)

        # scan rate
        scan = None
        if not r['ScanRate']:
            E('ScanRate is empty (whole seconds)')
        else:
            try:
                scan = to_int(r['ScanRate'])
                if scan < 1:
                    E('ScanRate must be at least 1 second')
            except ValueError:
                E('ScanRate "%s" must be whole seconds' % r['ScanRate'])

        if len(errors) > n_err:
            continue
        discrete = fmt == 'DISCRETE'
        prefix, tmpl = AREAS[area]
        tags.append({
            'row': rn, 'device': dev, 'area': area, 'addr': addr_out, 'fmt': fmt, 'length': length,
            'start': start,
            'cells': [full, 0.0, '0' if discrete else '0.0', 1.0 if discrete else 0.0,
                      '%s%04d' % (prefix, addr_out), 0.0, 0.0 if discrete else 100.0, rev,
                      r['Description'], '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', rw, tmpl,
                      str(length), str(FORMATS[fmt][0]), str(start), str(scan)],
        })

    # register overlap (information only) and address-space tail
    by_area = {}
    for t in tags:
        regs = max(1, math.ceil(t['length'] / 16)) if t['area'] in ('HR', 'IR') else 1
        last = t['addr'] + regs - 1
        if last > 9999:
            confirms.append((t['row'], '%s spans registers %d..%d, past 9999' % (t['cells'][0], t['addr'], last)))
        bit = t['start'] if (t['fmt'] == 'DISCRETE' and t['area'] in ('HR', 'IR')) else None
        for reg in range(t['addr'], last + 1):
            by_area.setdefault((t['device'], t['area'], reg), []).append((t, bit))
    # union-find: tags sharing a register (bits on different positions of one register do not clash)
    parent = {t['row']: t['row'] for t in tags}
    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x
    clash_regs = {}
    for (dev, area, reg), lst in by_area.items():
        for i in range(len(lst)):
            for j in range(i + 1, len(lst)):
                (a, ba), (b, bb) = lst[i], lst[j]
                if ba is not None and bb is not None and ba != bb:
                    continue
                parent[find(a['row'])] = find(b['row'])
                clash_regs.setdefault((a['row'], b['row']), (area, set()))[1].add(reg)
    groups = {}
    for (ra, rb), (area, regs) in clash_regs.items():
        g = groups.setdefault(find(ra), {'rows': set(), 'regs': set(), 'area': area})
        g['rows'].update((ra, rb))
        g['regs'].update(regs)
    by_row = {t['row']: t for t in tags}
    for g in groups.values():
        members = [by_row[r] for r in sorted(g['rows'])]
        lo, hi = min(g['regs']), max(g['regs'])
        span = '%s%04d' % (AREAS[g['area']][0], lo) + ('' if lo == hi else '..%s%04d' % (AREAS[g['area']][0], hi))
        sig = lambda t: tuple(c for k, c in enumerate(t['cells']) if k not in (0, 8))
        dup = len(members) != len({sig(t) for t in members})
        notes.append((members[0]['row'], '%d tags share %s: %s%s' % (
            len(members), span, ', '.join('%s (row %d)' % (t['cells'][0], t['row']) for t in members),
            ' - includes identical definitions' if dup else '')))
    return tags, errors, confirms, notes


# ---------------------------------------------------------------- BoardIO
def boardio_builtin():
    rows = []
    def row(n, dt, span_hi, rw, length):
        return ['BoardIO:' + n, 0.0, '0', dt, n, 0.0, span_hi, 'false', '', '0', 0.0, 0.0, 0.0, 0.0,
                'false', 'false', 'false', rw, 0.0, length, '0', '0', '1']
    rows += [row('AI.%d' % i, 0.0, 0.0, '1', '16') for i in range(8)]
    rows += [row('DI.%d' % i, 1.0, FLT_MAX, '1', '1') for i in range(8)]
    rows += [row('DO.%d' % i, 1.0, FLT_MAX, '3', '1') for i in range(4)]
    return rows


def boardio_from(path):
    import xlrd
    s = xlrd.open_workbook(path).sheet_by_name('BoardIO')
    if [s.cell_value(0, c) for c in range(s.ncols)] != HEADERS:
        sys.exit('BoardIO sheet in %s does not have the expected headers' % path)
    out = []
    for r in range(1, s.nrows):
        out.append([float(c.value) if c.ctype == xlrd.XL_CELL_NUMBER else
                    ('' if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else str(c.value))
                    for c in s.row(r)])
    return out


# ---------------------------------------------------------------- xls writer
def write_xls(path, sheets):
    """sheets: list of (sheet_name, headers, rows, widths). Text cells stay text (even ''), numbers stay numbers."""
    import xlwt
    from xlwt.Cell import StrCell
    wb = xlwt.Workbook(encoding='utf-8')
    for name, headers, rows, widths in sheets:
        ws = wb.add_sheet(name)
        for c, w in enumerate(widths):
            ws.col(c).width = w
        for r, values in enumerate([headers] + rows):
            row = ws.row(r)
            for c, v in enumerate(values):
                if isinstance(v, str) and v == '':
                    row.write(c, '')                       # blank cell (sets bounds/style) ...
                    cells = row._Row__cells
                    cells[c] = StrCell(r, c, cells[c].xf_idx, wb.add_str(''))  # ... then make it text ''
                else:
                    row.write(c, v)
    wb.save(path)


def verify_xls(path, sheets):
    import xlrd
    wb = xlrd.open_workbook(path)
    problems = []
    if wb.sheet_names() != [s[0] for s in sheets]:
        problems.append('sheet order %s' % wb.sheet_names())
    for name, headers, rows, _ in sheets:
        s = wb.sheet_by_name(name)
        for r, values in enumerate([headers] + rows):
            for c, v in enumerate(values):
                cell = s.cell(r, c)
                want = xlrd.XL_CELL_TEXT if isinstance(v, str) else xlrd.XL_CELL_NUMBER
                if cell.ctype != want or cell.value != v:
                    problems.append('%s!R%dC%d: %r (type %d) != %r' % (name, r + 1, c + 1, cell.value, cell.ctype, v))
    return problems


# ---------------------------------------------------------------- input template / decode
def write_input_template(path, records=(), note=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tags'
    ws.append(IN_COLS)
    bold, fill = Font(bold=True), PatternFill('solid', fgColor='DDEBF7')
    for c in ws[1]:
        c.font, c.fill = bold, fill
    for col, w in zip('ABCDEFGHIJK', [14, 22, 13, 10, 13, 8, 8, 9, 10, 9, 60]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    for rec in records:
        ws.append([rec.get(c, '') for c in IN_COLS])
    n = max(2000, len(records) + 500)
    for col, values in (('C', 'CS,IS,IR,HR'), ('E', ','.join(FORMATS)), ('F', '1,8,16,32,64'),
                        ('G', 'R,RW'), ('I', 'TRUE,FALSE')):
        dv = DataValidation(type='list', formula1='"%s"' % values, allow_blank=True,
                            errorStyle='warning', showErrorMessage=True)
        ws.add_data_validation(dv)
        dv.add('%s2:%s%d' % (col, col, n))

    ref = wb.create_sheet('Reference')
    lines = [
        ['ADAM-3600 tag input - one row per tag. Generated .xls follows adam-3600_tags_template.xls.'],
        [note or 'Address base (0 or 1) is asked every time the import file is generated.'],
        [],
        ['Column', 'Required', 'Meaning'],
        ['Device', 'yes', 'I/O device name = sheet name = left side of "Device:Tag". No ":"; not BoardIO/User Tag/Calc Tag.'],
        ['Tag', 'yes', 'Tag name = right side of "Device:Tag". Full "Device:Tag" must be <= 26 characters.'],
        ['RegisterType', 'yes', 'CS (FC01, prefix 0), IS (FC02, prefix 1), IR (FC04, prefix 3), HR (FC03, prefix 4).'],
        ['Address', 'yes', 'Register number inside the area, without the prefix. Base-0 or base-1 - asked at generation. Result must be 1..9999.'],
        ['Format', 'yes', 'See formats below. A width suffix (INT32, FLOAT32, BCD16) can replace the Length column.'],
        ['Length', 'analog', 'Bits read from Modbus: 8, 16, 32 or 64. Discrete = 1 (may be left blank).'],
        ['Access', 'yes', 'R (readWrite 1) or RW (readWrite 3). IS and IR are always R.'],
        ['StartBit', 'discrete on HR/IR', 'Bit index 0..15 inside the register. Blank = 0 elsewhere.'],
        ['ActiveLow', 'no', 'TRUE = signal reversed (active low). Blank = FALSE.'],
        ['ScanRate', 'yes', 'Whole seconds (>= 1).'],
        ['Description', 'no', 'Free text notes.'],
        [],
        ['Format', 'conversionType', 'Meaning', 'Usual length'],
    ]
    for f, (code, meaning) in FORMATS.items():
        lines.append([f, code, meaning, {'DISCRETE': 1}.get(f, EXPECTED_LEN.get(f, '8/16/32/64'))])
    for l in lines:
        ref.append(l)
    for c in ref[4] + ref[17]:
        c.font = bold
    for col, w in zip('ABCD', [14, 16, 95, 14]):
        ref.column_dimensions[col].width = w
    for row in ref.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(path)


def decode(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    recs, issues = [], []
    for s in wb.sheets():
        if s.name.lower() in RESERVED_SHEETS:
            if s.name != 'BoardIO' and s.nrows > 1:
                issues.append('%s has %d rows - not decoded (not supported yet)' % (s.name, s.nrows - 1))
            continue
        if s.nrows == 0 or [s.cell_value(0, c) for c in range(s.ncols)] != HEADERS:
            issues.append('sheet %s skipped - headers differ from the tag layout' % s.name)
            continue
        for r in range(1, s.nrows):
            v = dict(zip(HEADERS, s.row_values(r)))
            name = str(v['name'])
            dev, _, tag = name.partition(':')
            addr = str(v['address'])
            area = PREFIX_TO_AREA.get(addr[:1])
            conv = int(float(v['conversionType']))
            fmt = 'DISCRETE' if float(v['dataType']) == 1 else CONV_TO_ANALOG.get(conv)
            if fmt == 'DISCRETE' and conv != 21:
                fmt = None
            if not area or not fmt or len(addr) != 5:
                issues.append('%s row %d (%s): address %s / conversionType %s not decodable' % (s.name, r + 1, name, addr, conv))
                continue
            discrete = fmt == 'DISCRETE'
            defaults = {'Scale.spanLow': 0, 'Scale.spanHigh': 0 if discrete else 100, 'scalingType': '0',
                        'Scale.scale': 0, 'Scale.bias': 0, 'Scale.outputLow': 0, 'Scale.outputHigh': 0,
                        'Scale.clampToLow': 'false', 'Scale.clampToHigh': 'false', 'Scale.clampToZero': 'false',
                        'type': 0}
            for k, d in defaults.items():
                if cell_str(v[k]) != cell_str(d if not isinstance(d, int) else float(d)):
                    issues.append('%s row %d (%s): %s=%r is not representable in the input (will become %r)'
                                  % (s.name, r + 1, name, k, v[k], d))
            recs.append({'Device': dev, 'Tag': tag, 'RegisterType': area, 'Address': int(addr[1:]),
                         'Format': fmt, 'Length': int(float(v['length'])),
                         'Access': 'RW' if cell_str(v['readWrite']) == '3' else 'R',
                         'StartBit': int(float(v['startBit'])),
                         'ActiveLow': 'TRUE' if str(v['signalReverse']).lower() == 'true' else 'FALSE',
                         'ScanRate': int(float(v['scanRate'])), 'Description': v['description']})
    return recs, issues


# ---------------------------------------------------------------- self-test
SAMPLE_INPUT = [  # VA_Temp sheet of adam-3600_tags_template.xls in canonical form (base-1)
    ('VA_Temp', 'Discrete_RW_CS', 'CS', '1', 'DISCRETE', '1', 'RW', '0', 'FALSE', '1', 'Discrete Read/Write Coil Status'),
    ('VA_Temp', 'Discrete_R_IS', 'IS', '1', 'DISCRETE', '1', 'R', '0', 'FALSE', '1', 'Discrete Read-Only Input Status (IS)'),
    ('VA_Temp', 'Discrete_RW_HR', 'HR', '1', 'DISCRETE', '1', 'RW', '5', 'FALSE', '1', 'Discrete Read/Write on Holding Register (HR), bit index 5 active high'),
    ('VA_Temp', 'Discrete_R_IR', 'IR', '1', 'DISCRETE', '1', 'R', '3', 'TRUE', '1', 'Discrete Read/Write on Input Register (IR), bit index 3 active low'),
    ('VA_Temp', 'Analog_UINT16_HR', 'HR', '1', 'UINT', '16', 'RW', '0', 'FALSE', '1', 'Unsigned Integer 16 bit on Holding Register, starts at bit 0 enabling Read/Write'),
    ('VA_Temp', 'Analog_INT16_HR', 'HR', '1', 'INT', '16', 'RW', '0', 'FALSE', '5', 'Integer 16 bit starting at bit 0 with address 1, scan rate 5 sec and read only'),
    ('VA_Temp', 'Analog_Double64_IR', 'IR', '1', 'DOUBLE', '64', 'R', '0', 'FALSE', '1', 'Analog Double Percision on IR register, Read-only'),
    ('VA_Temp', 'Analog_Float32_HR', 'HR', '1', 'FLOAT', '32', 'RW', '0', 'FALSE', '1', 'Read/Write HR IEEE 32 bit floating point on address 1 with scan rate of 1, start at bit 0\r\n'),
    ('VA_Temp', 'Analog_RFloat32_HR', 'HR', '5', 'RFLOAT', '32', 'RW', '0', 'FALSE', '1', 'Reverse Word Floating Point 32 bit wide start at address 5, with read/write enabled, scan rate 1'),
    ('VA_Temp', 'Analog_HIWORD_INT', 'HR', '1', 'HIWORD_INT', '16', 'RW', '0', 'FALSE', '1', 'High Word Integer 16 bit on HR register starting at bit 0 and address 1 with scan rate of 1 and read/write enabled'),
    ('VA_Temp', 'Analog_HIWORD_UINT', 'HR', '1001', 'HIWORD_UINT', '16', 'RW', '0', 'FALSE', '1', 'Unsigned Integer High Word starting at address 1001, and start bit 0, with scan rate of 1 and read/write enabled'),
    ('VA_Temp', 'Analog_BCD8_IR', 'IR', '1', 'BCD', '8', 'R', '0', 'FALSE', '1', '8 bit BCD on IR register start at address 1 and bit 0, read-only, and scan rate 1'),
    ('VA_Temp', 'Analog_INT16_SWAP', 'HR', '1', 'INT_SWAP', '16', 'RW', '0', 'FALSE', '1', 'Swap Byte Singed Integer 16 bit start at address 1 and bit 0, scan rate 1, read/write on HR register'),
    ('VA_Temp', 'Analog_UINT16_SWAP', 'HR', '1', 'UINT_SWAP', '16', 'RW', '0', 'FALSE', '1', 'HR register 16 bit unsigned integer swap byte start at address 1 bit 0 and scan rate 1 with read/write'),
    ('VA_Temp', 'Analog_Double_BE', 'HR', '1', 'DOUBLE_BE', '64', 'R', '0', 'FALSE', '1', 'Double 64-bit big endian on HR register start at address 1 bit 0 and scan rate 1 with read-only '),
    ('VA_Temp', 'Analog_Double_LE', 'IR', '1', 'DOUBLE_LE', '64', 'R', '0', 'FALSE', '1', 'IR read-only start at address 1 bit 0 read-only with 64 bit little endian double'),
]
SAMPLE_EXPECTED = [  # VA_Temp rows exactly as in the template
    ['VA_Temp:Discrete_RW_CS', 0.0, '0', 1.0, '00001', 0.0, 0.0, 'false', 'Discrete Read/Write Coil Status', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 0.0, '1', '21', '0', '1'],
    ['VA_Temp:Discrete_R_IS', 0.0, '0', 1.0, '10001', 0.0, 0.0, 'false', 'Discrete Read-Only Input Status (IS)', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '1', 1.0, '1', '21', '0', '1'],
    ['VA_Temp:Discrete_RW_HR', 0.0, '0', 1.0, '40001', 0.0, 0.0, 'false', 'Discrete Read/Write on Holding Register (HR), bit index 5 active high', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '1', '21', '5', '1'],
    ['VA_Temp:Discrete_R_IR', 0.0, '0', 1.0, '30001', 0.0, 0.0, 'true', 'Discrete Read/Write on Input Register (IR), bit index 3 active low', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '1', 2.0, '1', '21', '3', '1'],
    ['VA_Temp:Analog_UINT16_HR', 0.0, '0.0', 0.0, '40001', 0.0, 100.0, 'false', 'Unsigned Integer 16 bit on Holding Register, starts at bit 0 enabling Read/Write', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '16', '21', '0', '1'],
    ['VA_Temp:Analog_INT16_HR', 0.0, '0.0', 0.0, '40001', 0.0, 100.0, 'false', 'Integer 16 bit starting at bit 0 with address 1, scan rate 5 sec and read only', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '16', '22', '0', '5'],
    ['VA_Temp:Analog_Double64_IR', 0.0, '0.0', 0.0, '30001', 0.0, 100.0, 'false', 'Analog Double Percision on IR register, Read-only', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '1', 2.0, '64', '61', '0', '1'],
    ['VA_Temp:Analog_Float32_HR', 0.0, '0.0', 0.0, '40001', 0.0, 100.0, 'false', 'Read/Write HR IEEE 32 bit floating point on address 1 with scan rate of 1, start at bit 0\r\n', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '32', '25', '0', '1'],
    ['VA_Temp:Analog_RFloat32_HR', 0.0, '0.0', 0.0, '40005', 0.0, 100.0, 'false', 'Reverse Word Floating Point 32 bit wide start at address 5, with read/write enabled, scan rate 1', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '32', '26', '0', '1'],
    ['VA_Temp:Analog_HIWORD_INT', 0.0, '0.0', 0.0, '40001', 0.0, 100.0, 'false', 'High Word Integer 16 bit on HR register starting at bit 0 and address 1 with scan rate of 1 and read/write enabled', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '16', '16', '0', '1'],
    ['VA_Temp:Analog_HIWORD_UINT', 0.0, '0.0', 0.0, '41001', 0.0, 100.0, 'false', 'Unsigned Integer High Word starting at address 1001, and start bit 0, with scan rate of 1 and read/write enabled', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '16', '17', '0', '1'],
    ['VA_Temp:Analog_BCD8_IR', 0.0, '0.0', 0.0, '30001', 0.0, 100.0, 'false', '8 bit BCD on IR register start at address 1 and bit 0, read-only, and scan rate 1', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '1', 2.0, '8', '35', '0', '1'],
    ['VA_Temp:Analog_INT16_SWAP', 0.0, '0.0', 0.0, '40001', 0.0, 100.0, 'false', 'Swap Byte Singed Integer 16 bit start at address 1 and bit 0, scan rate 1, read/write on HR register', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '16', '2', '0', '1'],
    ['VA_Temp:Analog_UINT16_SWAP', 0.0, '0.0', 0.0, '40001', 0.0, 100.0, 'false', 'HR register 16 bit unsigned integer swap byte start at address 1 bit 0 and scan rate 1 with read/write', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '3', 3.0, '16', '1', '0', '1'],
    ['VA_Temp:Analog_Double_BE', 0.0, '0.0', 0.0, '40001', 0.0, 100.0, 'false', 'Double 64-bit big endian on HR register start at address 1 bit 0 and scan rate 1 with read-only ', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '1', 3.0, '64', '150', '0', '1'],
    ['VA_Temp:Analog_Double_LE', 0.0, '0.0', 0.0, '30001', 0.0, 100.0, 'false', 'IR read-only start at address 1 bit 0 read-only with 64 bit little endian double', '0', 0.0, 0.0, 0.0, 0.0, 'false', 'false', 'false', '1', 2.0, '64', '151', '0', '1'],
]


def selftest():
    """Rebuild the template's VA_Temp sheet from SAMPLE_INPUT and compare every cell (value + text/number type)."""
    import tempfile
    fails = []
    for base, shift in ((1, 0), (0, -1)):
        rows = []
        for i, rec in enumerate(SAMPLE_INPUT):
            d = dict(zip(IN_COLS, rec))
            d['Address'] = str(int(d['Address']) + shift)
            rows.append((i + 2, d))
        tags, errors, confirms, _ = build_tags(rows, base)
        if errors or confirms:
            fails.append('base-%d: unexpected errors/confirmations %s' % (base, errors + confirms))
            continue
        got = [t['cells'] for t in tags]
        for r, (g, e) in enumerate(zip(got, SAMPLE_EXPECTED)):
            for c, (gv, ev) in enumerate(zip(g, e)):
                if gv != ev or type(gv) is not type(ev):
                    fails.append('base-%d row %d %s: %r != %r' % (base, r + 2, HEADERS[c], gv, ev))
        if len(got) != len(SAMPLE_EXPECTED):
            fails.append('base-%d: %d rows != %d' % (base, len(got), len(SAMPLE_EXPECTED)))
    fd, tmp = tempfile.mkstemp(suffix='.xls')
    os.close(fd)
    sheets = [('BoardIO', HEADERS, boardio_builtin(), WIDTHS['BoardIO']),
              ('VA_Temp', HEADERS, SAMPLE_EXPECTED, WIDTHS['device']),
              ('User Tag', USER_TAG_HEADERS, [], WIDTHS['User Tag']),
              ('Calc Tag', CALC_TAG_HEADERS, [], WIDTHS['Calc Tag'])]
    write_xls(tmp, sheets)
    fails += verify_xls(tmp, sheets)
    os.remove(tmp)
    if fails:
        print('SELFTEST FAILED - do not use this script until fixed:', *fails[:30], sep='\n  ')
        return 1
    print('selftest OK - 16 template tags reproduced exactly (base-1 and base-0), .xls write/read verified')
    return 0


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    t = sub.add_parser('template')
    t.add_argument('out')
    b = sub.add_parser('build')
    b.add_argument('input')
    b.add_argument('--base', type=int, choices=(0, 1), required=True,
                   help='0: input address 0 = Modbus offset 0 (written +1). 1: input address 1 = offset 0 (written as is)')
    b.add_argument('-o', '--out', required=True)
    b.add_argument('--sheet')
    b.add_argument('--boardio', help='copy the BoardIO sheet from this existing ADAM-3600 .xls instead of the built-in one')
    b.add_argument('--confirm', action='store_true', help='user approved every confirmation item')
    sub.add_parser('selftest')
    d = sub.add_parser('decode')
    d.add_argument('export')
    d.add_argument('-o', '--out', required=True)
    a = ap.parse_args()

    if a.cmd == 'selftest':
        return selftest()

    if a.cmd == 'template':
        write_input_template(a.out)
        print('Wrote', a.out)
        return 0

    if a.cmd == 'decode':
        recs, issues = decode(a.export)
        write_input_template(a.out, recs, note='Decoded from %s. Addresses are base-1 (as stored in the export) - '
                             'answer base-1 when generating from this file.' % os.path.basename(a.export))
        print('Decoded %d tags -> %s (addresses base-1)' % (len(recs), a.out))
        for i in issues:
            print('  NOTE:', i)
        return 0

    rows = read_table(a.input, a.sheet)
    tags, errors, confirms, notes = build_tags(rows, a.base)
    print('Input rows: %d | valid: %d | errors: %d | need confirmation: %d | overlap notes: %d'
          % (len(rows), len(tags), len(errors), len(confirms), len(notes)))
    for label, items in (('ERROR', errors), ('CONFIRM', confirms), ('NOTE', notes)):
        for rn, m in sorted(items):
            print('  %-7s row %d: %s' % (label, rn, m))
    if errors:
        print('Nothing written - fix the errors above.')
        return 1
    if confirms and not a.confirm:
        print('Nothing written - get the user\'s approval for the CONFIRM items, then rerun with --confirm.')
        return 2

    devices = []
    for t_ in tags:
        if t_['device'] not in devices:
            devices.append(t_['device'])
    sheets = [('BoardIO', HEADERS, boardio_from(a.boardio) if a.boardio else boardio_builtin(), WIDTHS['BoardIO'])]
    for dv in devices:
        drows = [t_['cells'] for t_ in tags if t_['device'] == dv]
        if len(drows) > 65534:
            print('ERROR: device %s has %d tags - more than one .xls sheet can hold' % (dv, len(drows)))
            return 1
        sheets.append((dv, HEADERS, drows, WIDTHS['device']))
    sheets += [('User Tag', USER_TAG_HEADERS, [], WIDTHS['User Tag']),
               ('Calc Tag', CALC_TAG_HEADERS, [], WIDTHS['Calc Tag'])]
    write_xls(a.out, sheets)
    problems = verify_xls(a.out, sheets)
    if problems:
        print('VERIFY FAILED:', *problems[:20], sep='\n  ')
        return 1
    print('Wrote %s (address base-%d input) - re-read and verified.' % (a.out, a.base))
    for dv in devices:
        dt = [t_ for t_ in tags if t_['device'] == dv]
        counts = {}
        for t_ in dt:
            k = '%s/%d' % (t_['fmt'], t_['length'])
            counts[k] = counts.get(k, 0) + 1
        print('  %s: %d tags  (%s)' % (dv, len(dt), ', '.join('%s x%d' % kv for kv in sorted(counts.items()))))
    return 0


if __name__ == '__main__':
    sys.exit(main())
